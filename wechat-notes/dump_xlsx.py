import os
import sys
import openpyxl

for path in sys.argv[1:]:
    lines = []
    lines.append("=" * 100)
    lines.append("FILE: " + path)
    wb = openpyxl.load_workbook(path, data_only=True)
    for ws in wb.worksheets:
        lines.append(f"\n--- sheet: {ws.title}  dims={ws.dimensions} "
                     f"max_row={ws.max_row} max_col={ws.max_column} ---")
        if ws.merged_cells.ranges:
            lines.append("merged: " + ", ".join(str(r) for r in ws.merged_cells.ranges))
        for row in ws.iter_rows():
            cells = []
            for c in row:
                v = c.value
                if v is None:
                    continue
                s = str(v).replace("\r\n", "\n").replace("\n", "\\n")
                cells.append(f"{c.coordinate}={s}")
            if cells:
                lines.append("  " + " | ".join(cells))
    base = os.path.splitext(os.path.basename(path))[0]
    out = os.path.join(os.path.dirname(os.path.abspath(__file__)), base + ".dump.txt")
    with open(out, "w", encoding="utf-8") as f:
        f.write("\n".join(lines) + "\n")
    print("wrote", out)
