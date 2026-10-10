# -*- coding: utf-8 -*-
"""把 normalized.json 生成一份完整的 Word 文档（高考英语 3500 词）"""
import json
import os
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

SRC = 'G:/dsh/gaokao-3500/_source/normalized.json'
OUT = 'G:/dsh/gaokao-3500/文档/高考3500词-完整文档.docx'
os.makedirs(os.path.dirname(OUT), exist_ok=True)

data = json.load(open(SRC, encoding='utf-8'))
stats = data['stats']


def set_cjk(style, ascii_font='Times New Roman', cjk_font='宋体'):
    rpr = style.element.get_or_add_rPr()
    rf = rpr.find(qn('w:rFonts'))
    if rf is None:
        rf = OxmlElement('w:rFonts')
        rpr.append(rf)
    rf.set(qn('w:ascii'), ascii_font)
    rf.set(qn('w:hAnsi'), ascii_font)
    rf.set(qn('w:eastAsia'), cjk_font)


def repeat_header(row):
    trpr = row._tr.get_or_add_trPr()
    el = OxmlElement('w:tblHeader')
    el.set(qn('w:val'), 'true')
    trpr.append(el)


def add_page_number(paragraph):
    run = paragraph.add_run()
    b = OxmlElement('w:fldChar'); b.set(qn('w:fldCharType'), 'begin')
    i = OxmlElement('w:instrText'); i.set(qn('xml:space'), 'preserve'); i.text = 'PAGE'
    e = OxmlElement('w:fldChar'); e.set(qn('w:fldCharType'), 'end')
    run._r.append(b); run._r.append(i); run._r.append(e)


doc = Document()

# 页面与样式
sec = doc.sections[0]
sec.top_margin = Cm(1.6); sec.bottom_margin = Cm(1.6)
sec.left_margin = Cm(1.5); sec.right_margin = Cm(1.5)

set_cjk(doc.styles['Normal'])
doc.styles['Normal'].font.size = Pt(9)
for name, size in (('Title', 22), ('Heading 1', 15), ('Heading 2', 12), ('Heading 3', 10.5)):
    try:
        st = doc.styles[name]
        set_cjk(st, cjk_font='黑体')
        st.font.size = Pt(size)
        st.font.color.rgb = RGBColor(0x1F, 0x3A, 0x5F)
    except KeyError:
        pass

footer = sec.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
footer.add_run('高考英语 3500 词 · 第 ')
add_page_number(footer)
footer.add_run(' 页')


def para(text, size=9, bold=False, space_after=3, align=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(0)
    r = p.add_run(text)
    r.bold = bold
    r.font.size = Pt(size)
    if align:
        p.alignment = align
    return p


def make_table(headers, rows, widths):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = 'Table Grid'
    t.autofit = False
    hdr = t.rows[0].cells
    for i, h in enumerate(headers):
        p = hdr[i].paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(h)
        r.bold = True
        r.font.size = Pt(8.5)
    repeat_header(t.rows[0])
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            p = cells[i].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(str(v))
            r.font.size = Pt(8.5)
    for i, w in enumerate(widths):
        for row in t.rows:
            row.cells[i].width = Cm(w)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t


def trans_of(e):
    return '；'.join(e['trans']).replace('|', '／')


def flag_of(e):
    f = ''
    if e.get('zhenti'):
        f += '★'
    if e.get('yuedu'):
        f += '●'
    if e.get('mine'):
        f += '✎'
    return f


# ---------------- 封面与说明 ----------------
title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
tr = title.add_run('高考英语 3500 词总表')
tr.bold = True
tr.font.size = Pt(26)

sub = doc.add_paragraph()
sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
sr = sub.add_run('单词 %d 个　短语/合成词 %d 条　合计 %d 条' % (stats['words'], stats['phrases'], stats['total']))
sr.font.size = Pt(11)

para('')
para('标记图例', size=11, bold=True)
para('★ 高考真题核心高频词（%d 个，时间紧就先背这批）' % stats['zhenti'])
para('● 高考阅读高频词（%d 个）' % stats['yuedu'])
para('✎ 本人实测错词（%d 个，详见第四部分）' % stats['mine'])
para('无标记 = 大纲基础词，仍需掌握')
para('')
para('怎么用', size=11, bold=True)
para('1. 遮住“释义”列自测：2 秒内说出意思的算过，犹豫或说错的抄进自己的错词本。')
para('2. 抄下来的词按 D+1 / D+2 / D+4 / D+7 / D+15 复习，第一次复习务必在 24 小时内。')
para('3. 时间不够时，优先刷第一部分 ★ 真题核心高频词。')
para('4. 写作用第三部分短语，比只会单词更像“人写的英语”。')
para('')
para('资料来源：高考大纲词表（星火抗遗忘高考英语3500词，含 %d 单词 + %d 短语），音标与真题高频标记取自公开词库；'
     '带音标条目 %d 条。' % (stats['words'], stats['phrases'], stats['withPhone']), size=8)
doc.add_page_break()

# ---------------- 一、真题核心高频 ----------------
doc.add_heading('一、★ 真题核心高频 %d 词' % stats['zhenti'], level=1)
zrows = [[i + 1, e['word'], e['us'], trans_of(e)] for i, e in enumerate(data['zhead'])]
make_table(['#', '单词', '音标(美)', '释义'], zrows, [1.1, 2.9, 3.2, 10.6])

# ---------------- 二、A–Z 单词表 ----------------
doc.add_page_break()
doc.add_heading('二、A–Z 单词表（%d 个）' % stats['words'], level=1)
for g in data['groups']:
    doc.add_heading('%s（%d 个）' % (g['letter'], len(g['items'])), level=2)
    rows = [[e['no'], e['word'], e['us'], trans_of(e), flag_of(e)] for e in g['items']]
    make_table(['#', '单词', '音标(美)', '释义', '标记'], rows, [1.1, 2.9, 3.2, 9.7, 0.9])

# ---------------- 三、短语 ----------------
doc.add_page_break()
doc.add_heading('三、短语与合成词（%d 条）' % stats['phrases'], level=1)
for g in data['phraseGroups']:
    doc.add_heading('%s（%d 条）' % (g['letter'], len(g['items'])), level=2)
    rows = [[e['no'], e['word'], trans_of(e), flag_of(e)] for e in g['items']]
    make_table(['#', '短语', '释义', '标记'], rows, [1.1, 4.4, 11.4, 0.9])

# ---------------- 附录、真题高频补充词 ----------------
if data['extra']:
    doc.add_page_break()
    doc.add_heading('附、真题高频补充词（不在上表内的 %d 个）' % len(data['extra']), level=1)
    rows = [[i + 1, e['word'], e['us'], trans_of(e)] for i, e in enumerate(data['extra'])]
    make_table(['#', '单词', '音标(美)', '释义'], rows, [1.1, 2.9, 3.2, 10.6])

# ---------------- 四、我的错词本 ----------------
doc.add_page_break()
doc.add_heading('四、我的错词本（实测抓取 %d 个）' % len(data['mine']), level=1)
para('这些是逐题实测中真实写错/写不出的词，比任何“高频词表”都更该优先复习。', size=9)
rows = [[i + 1, m['display'], m['tip']] for i, m in enumerate(data['mine'])]
make_table(['#', '单词', '中文 / 辨析'], rows, [1.1, 4.6, 12.1])

doc.save(OUT)
print('已生成:', OUT)
print('大小: %.1f MB' % (os.path.getsize(OUT) / 1024 / 1024))
print('统计:', stats)
